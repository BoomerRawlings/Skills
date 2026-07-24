#!/usr/bin/env python3
"""Audit and apply reviewed grayscale-safe visual edits."""

from __future__ import annotations

import argparse
import io
import json
import math
import re
import shutil
import sys
import tempfile
import zipfile
from collections import Counter
from itertools import combinations
from pathlib import Path

try:
    import fitz
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt
    from docx.text.paragraph import Paragraph
    from PIL import Image, ImageChops, ImageDraw, ImageEnhance, ImageFont, ImageOps, ImageStat
    from reportlab.lib.pagesizes import LETTER, landscape
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas
except ImportError as exc:
    raise SystemExit(f"Missing required runtime dependency: {exc.name}") from exc


class BWError(RuntimeError):
    def __init__(self, reason: str, resolution: str, prompt: str):
        super().__init__(reason)
        self.reason = reason
        self.resolution = resolution
        self.prompt = prompt


def blocked(reason: str, resolution: str, prompt: str) -> None:
    raise BWError(reason, resolution, prompt)


def norm(value: str) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def parse_color(value: str) -> tuple[int, int, int]:
    match = re.fullmatch(r"#?([0-9a-fA-F]{6})", norm(value))
    if not match:
        blocked(f"Invalid color '{value}'.", "Use a six-digit hex color.", "Correct the color? yes/no")
    raw = match.group(1)
    return tuple(int(raw[index:index + 2], 16) for index in (0, 2, 4))


def color_hex(color: tuple[int, int, int]) -> str:
    return "#%02X%02X%02X" % color


def linear_channel(value: int) -> float:
    channel = value / 255
    return channel / 12.92 if channel <= 0.04045 else ((channel + 0.055) / 1.055) ** 2.4


def luminance(color: tuple[int, int, int]) -> float:
    r, g, b = (linear_channel(value) for value in color)
    return 0.2126 * r + 0.7152 * g + 0.0722 * b


def contrast_ratio(first: tuple[int, int, int], second: tuple[int, int, int]) -> float:
    high, low = sorted((luminance(first), luminance(second)), reverse=True)
    return (high + 0.05) / (low + 0.05)


def lab(color: tuple[int, int, int]) -> tuple[float, float, float]:
    r, g, b = (linear_channel(value) for value in color)
    x = (0.4124 * r + 0.3576 * g + 0.1805 * b) / 0.95047
    y = (0.2126 * r + 0.7152 * g + 0.0722 * b)
    z = (0.0193 * r + 0.1192 * g + 0.9505 * b) / 1.08883

    def pivot(value: float) -> float:
        return value ** (1 / 3) if value > 0.008856 else 7.787 * value + 16 / 116

    fx, fy, fz = pivot(x), pivot(y), pivot(z)
    return 116 * fy - 16, 500 * (fx - fy), 200 * (fy - fz)


def delta_e(first: tuple[int, int, int], second: tuple[int, int, int]) -> float:
    left, right = lab(first), lab(second)
    return math.sqrt(sum((a - b) ** 2 for a, b in zip(left, right)))


def audit_image(image: Image.Image) -> dict:
    sample = image.convert("RGB")
    sample.thumbnail((900, 900), Image.Resampling.LANCZOS)
    quantized = sample.quantize(colors=12, method=Image.Quantize.MEDIANCUT)
    width, height = quantized.size
    pixels = list(quantized.get_flattened_data())
    counts = Counter(pixels)
    palette = quantized.getpalette() or []

    def palette_color(index: int) -> tuple[int, int, int]:
        start = index * 3
        return tuple(palette[start:start + 3])

    adjacency: Counter[tuple[int, int]] = Counter()
    for y in range(height):
        row = y * width
        for x in range(width):
            current = pixels[row + x]
            if x + 1 < width:
                other = pixels[row + x + 1]
                if current != other:
                    adjacency[tuple(sorted((current, other)))] += 1
            if y + 1 < height:
                other = pixels[row + width + x]
                if current != other:
                    adjacency[tuple(sorted((current, other)))] += 1
    minimum_adjacency = max(8, int(width * height * 0.0001))
    pixel_count = max(1, width * height)
    candidates: dict[tuple[int, int], dict] = {}
    for (left_index, right_index), adjacent_pixels in adjacency.items():
        left_fraction = counts[left_index] / pixel_count
        right_fraction = counts[right_index] / pixel_count
        if adjacent_pixels >= minimum_adjacency and min(left_fraction, right_fraction) >= 0.001:
            candidates[(left_index, right_index)] = {
                "relationship": "adjacent",
                "adjacent_sample_edges": adjacent_pixels,
            }
    major_indexes = [index for index, count in counts.items() if count / pixel_count >= 0.02]
    for left_index, right_index in combinations(sorted(major_indexes), 2):
        candidates.setdefault((left_index, right_index), {
            "relationship": "co-occurring",
            "adjacent_sample_edges": 0,
        })

    conflicts = []
    for (left_index, right_index), relationship in candidates.items():
        left, right = palette_color(left_index), palette_color(right_index)
        ratio = contrast_ratio(left, right)
        difference = delta_e(left, right)
        if ratio < 3 and difference >= 10:
            conflicts.append({
                "color_a": color_hex(left),
                "color_b": color_hex(right),
                "grayscale_contrast": round(ratio, 2),
                "color_difference_delta_e76": round(difference, 1),
                **relationship,
            })
    conflicts.sort(key=lambda item: (item["relationship"] != "adjacent", item["grayscale_contrast"]))
    dominant = [
        {"color": color_hex(palette_color(index)), "sample_fraction": round(count / max(1, len(pixels)), 4)}
        for index, count in counts.most_common(12)
    ]
    return {"width": image.width, "height": image.height, "conflicts": conflicts, "dominant_colors": dominant}


def laser_preview(image: Image.Image) -> Image.Image:
    gray = ImageOps.grayscale(image)
    gray = ImageOps.autocontrast(gray, cutoff=1)
    return gray.point(lambda value: min(255, int(round(value / 17)) * 17))


def save_previews(image: Image.Image, work_dir: Path, stem: str) -> dict:
    color_path = work_dir / f"{stem}-color.png"
    gray_path = work_dir / f"{stem}-gray.png"
    laser_path = work_dir / f"{stem}-laser.png"
    image.convert("RGB").save(color_path)
    ImageOps.grayscale(image).save(gray_path)
    laser_preview(image).save(laser_path)
    return {"color": str(color_path.resolve()), "grayscale": str(gray_path.resolve()), "laser": str(laser_path.resolve())}


def image_from_pixmap(pixmap: fitz.Pixmap) -> Image.Image:
    return Image.open(io.BytesIO(pixmap.tobytes("png"))).convert("RGB")


def pdf_audit(source: Path, work_dir: Path) -> list[dict]:
    doc = fitz.open(source)
    targets = []
    for page_index, page in enumerate(doc):
        full = image_from_pixmap(page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False))
        stem = f"page-{page_index + 1:04d}"
        full_report = audit_image(full)
        full_report.update({"id": f"page-{page_index + 1}-full", "kind": "pdf-page", "page": page_index + 1, "bbox": list(page.rect), "previews": save_previews(full, work_dir, stem)})
        targets.append(full_report)
        seen: set[tuple[float, float, float, float]] = set()
        image_number = 0
        for image_info in page.get_images(full=True):
            xref = image_info[0]
            for rect in page.get_image_rects(xref):
                key = tuple(round(value, 1) for value in rect)
                if key in seen or rect.width < 20 or rect.height < 20:
                    continue
                seen.add(key)
                image_number += 1
                region = image_from_pixmap(page.get_pixmap(matrix=fitz.Matrix(2.5, 2.5), clip=rect, alpha=False))
                region_stem = f"page-{page_index + 1:04d}-image-{image_number:02d}"
                report = audit_image(region)
                report.update({"id": f"page-{page_index + 1}-image-{image_number}", "kind": "pdf-image-region", "page": page_index + 1, "bbox": list(rect), "previews": save_previews(region, work_dir, region_stem)})
                targets.append(report)
    doc.close()
    return targets


def docx_occurrences(source: Path) -> list[dict]:
    document = Document(source)
    occurrences = []
    for paragraph_index, paragraph in enumerate(document.paragraphs, start=1):
        for blip in paragraph._p.xpath(".//a:blip"):
            relationship_id = blip.get(qn("r:embed"))
            if relationship_id and relationship_id in document.part.rels:
                target = document.part.rels[relationship_id].target_ref.replace("\\", "/")
                occurrences.append({"paragraph": paragraph_index, "media": f"word/{target}" if not target.startswith("word/") else target})
    return occurrences


def docx_audit(source: Path, work_dir: Path) -> tuple[list[dict], list[dict]]:
    targets = []
    with zipfile.ZipFile(source) as archive:
        for member in sorted(name for name in archive.namelist() if name.startswith("word/media/")):
            try:
                image = Image.open(io.BytesIO(archive.read(member))).convert("RGB")
            except Exception:
                continue
            stem = re.sub(r"[^a-zA-Z0-9_-]+", "-", Path(member).stem)
            report = audit_image(image)
            report.update({"id": member, "kind": "docx-media", "media": member, "previews": save_previews(image, work_dir, stem)})
            targets.append(report)
    return targets, docx_occurrences(source)


def image_audit(source: Path, work_dir: Path) -> list[dict]:
    image = Image.open(source).convert("RGB")
    report = audit_image(image)
    report.update({"id": "image-1", "kind": "standalone-image", "previews": save_previews(image, work_dir, "image-1")})
    return [report]


def audit(source: Path, work_dir: Path) -> dict:
    if work_dir.exists() and any(work_dir.iterdir()):
        blocked(f"Audit directory is not empty: {work_dir}", "Choose an empty audit directory.", "Use a new audit directory? yes/no")
    work_dir.mkdir(parents=True, exist_ok=True)
    suffix = source.suffix.lower()
    occurrences = []
    if suffix == ".pdf":
        input_type, targets = "pdf", pdf_audit(source, work_dir)
    elif suffix == ".docx":
        input_type = "docx"
        targets, occurrences = docx_audit(source, work_dir)
    elif suffix in (".png", ".jpg", ".jpeg"):
        input_type, targets = "image", image_audit(source, work_dir)
    else:
        blocked(f"Unsupported input format: {suffix or 'none'}.", "Use PDF, DOCX, PNG, or JPG.", "Choose a supported input? yes/no")
    return {
        "input": str(source.resolve()),
        "input_type": input_type,
        "thresholds": {"meaningful_graphics": "3:1 or redundant distinguisher", "normal_text": "4.5:1", "large_text": "3:1"},
        "warning": "Conflict pairs are visual candidates. Confirm semantic meaning from the source before editing.",
        "targets": targets,
        "docx_image_occurrences": occurrences,
        "conflict_count": sum(len(target["conflicts"]) for target in targets),
        "suggested_spec": {"edits": []},
    }


def color_mask(image: Image.Image, color: tuple[int, int, int], tolerance: int) -> Image.Image:
    tolerance_sq = max(0, tolerance) ** 2
    data = []
    for pixel in image.convert("RGB").get_flattened_data():
        distance = sum((pixel[index] - color[index]) ** 2 for index in range(3))
        data.append(255 if distance <= tolerance_sq else 0)
    mask = Image.new("L", image.size)
    mask.putdata(data)
    return mask


def pattern_mask(size: tuple[int, int], style: str, spacing: int, width: int) -> Image.Image:
    spacing = max(4, spacing)
    width = max(1, width)
    result = Image.new("L", size, 0)
    draw = ImageDraw.Draw(result)
    image_width, image_height = size
    if style in ("diagonal", "crosshatch"):
        for offset in range(-image_height, image_width + image_height, spacing):
            draw.line((offset, image_height, offset + image_height, 0), fill=255, width=width)
    if style in ("backward-diagonal", "crosshatch"):
        for offset in range(-image_height, image_width + image_height, spacing):
            draw.line((offset, 0, offset + image_height, image_height), fill=255, width=width)
    if style == "dots":
        radius = max(1, width)
        for y in range(spacing // 2, image_height, spacing):
            for x in range(spacing // 2, image_width, spacing):
                draw.ellipse((x - radius, y - radius, x + radius, y + radius), fill=255)
    if style not in ("diagonal", "backward-diagonal", "crosshatch", "dots"):
        blocked(f"Unsupported pattern '{style}'.", "Use diagonal, backward-diagonal, crosshatch, or dots.", "Correct the pattern? yes/no")
    return result


def load_font(size: int) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    choices = [Path("C:/Windows/Fonts/arial.ttf"), Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")]
    for path in choices:
        if path.exists():
            return ImageFont.truetype(str(path), max(8, size))
    return ImageFont.load_default()


def validate_caption(caption: str) -> str:
    caption = norm(caption)
    if not caption:
        return ""
    sentences = [part for part in re.split(r"(?<=[.!?])\s+", caption) if norm(part)]
    if len(sentences) > 3:
        blocked("Image description exceeds three sentences.", "Reduce it to 1-3 factual sentences.", "Shorten the description? yes/no")
    return caption


def edit_image(image: Image.Image, edit: dict) -> Image.Image:
    original = image.convert("RGB")
    result = original.copy()
    filters = dict(edit.get("filters", {}))
    if filters.get("autocontrast"):
        result = ImageOps.autocontrast(result, cutoff=1)
    result = ImageEnhance.Contrast(result).enhance(float(filters.get("contrast", 1.0)))
    result = ImageEnhance.Sharpness(result).enhance(float(filters.get("sharpness", 1.0)))
    for recolor in edit.get("recolors", []):
        mask = color_mask(original, parse_color(recolor["from"]), int(recolor.get("tolerance", 20)))
        result = Image.composite(Image.new("RGB", result.size, parse_color(recolor["to"])), result, mask)
    for pattern in edit.get("patterns", []):
        mask = color_mask(original, parse_color(pattern["color"]), int(pattern.get("tolerance", 20)))
        hatch = pattern_mask(result.size, pattern.get("style", "diagonal"), int(pattern.get("spacing", 10)), int(pattern.get("width", 1)))
        combined = ImageChops.multiply(mask, hatch)
        ink = parse_color(pattern.get("ink", "#000000"))
        result = Image.composite(Image.new("RGB", result.size, ink), result, combined)
    for outline in edit.get("outlines", []):
        width = max(1, int(outline.get("width", 3)))
        if width % 2 == 0:
            width += 1
        mask = color_mask(original, parse_color(outline["color"]), int(outline.get("tolerance", 20)))
        edge = ImageChops.subtract(mask.filter(ImageFilterMax(width)), mask.filter(ImageFilterMin(width)))
        ink = parse_color(outline.get("ink", "#000000"))
        result = Image.composite(Image.new("RGB", result.size, ink), result, edge)
    draw = ImageDraw.Draw(result)
    visual_scale = max(0.75, min(4.0, result.width / 1000))
    for marker in edit.get("markers", []):
        x = int(float(marker["x"]) * result.width)
        y = int(float(marker["y"]) * result.height)
        size = max(4, int(marker.get("size", 12) * visual_scale))
        box = (x - size, y - size, x + size, y + size)
        shape = marker.get("shape", "circle")
        if shape == "circle":
            draw.ellipse(box, fill="white", outline="black", width=2)
        elif shape == "square":
            draw.rectangle(box, fill="white", outline="black", width=2)
        elif shape == "triangle":
            draw.polygon(((x, y - size), (x - size, y + size), (x + size, y + size)), fill="white", outline="black")
        elif shape == "diamond":
            draw.polygon(((x, y - size), (x - size, y), (x, y + size), (x + size, y)), fill="white", outline="black")
        else:
            blocked(f"Unsupported marker '{shape}'.", "Use circle, square, triangle, or diamond.", "Correct the marker? yes/no")
    for label in edit.get("labels", []):
        text = norm(str(label.get("text", "")))
        if not text:
            blocked("A direct label has no text.", "Add factual label text.", "Correct the label? yes/no")
        x = int(float(label["x"]) * result.width)
        y = int(float(label["y"]) * result.height)
        font = load_font(max(7, int(label.get("size", 13) * visual_scale)))
        bbox = draw.textbbox((x, y), text, font=font)
        padded = (bbox[0] - 3, bbox[1] - 2, bbox[2] + 3, bbox[3] + 2)
        draw.rectangle(padded, fill="white", outline="black", width=1)
        draw.text((x, y), text, fill="black", font=font)
    return result


def ImageFilterMax(width: int):
    from PIL import ImageFilter
    return ImageFilter.MaxFilter(width)


def ImageFilterMin(width: int):
    from PIL import ImageFilter
    return ImageFilter.MinFilter(width)


def rect_from_target(target: dict, page: fitz.Page) -> fitz.Rect:
    values = target.get("bbox")
    if not isinstance(values, list) or len(values) != 4:
        blocked("A PDF edit is missing a four-number bbox.", "Specify [x0, y0, x1, y1] in page points.", "Specify the visual bounds? yes/no")
    rect = fitz.Rect(*(float(value) for value in values))
    if rect.is_empty or not page.rect.contains(rect):
        blocked(f"PDF visual bounds fall outside page {target.get('page')}.", "Correct the bbox.", "Correct the visual bounds? yes/no")
    return rect


def caption_space_blank(page: fitz.Page, rect: fitz.Rect) -> bool:
    image = image_from_pixmap(page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), clip=rect, alpha=False)).convert("L")
    stats = ImageStat.Stat(image)
    return stats.mean[0] > 242 and stats.stddev[0] < 8


def apply_pdf(source: Path, spec: dict, output: Path) -> None:
    doc = fitz.open(source)
    prepared: list[dict] = []
    by_page: dict[int, list[fitz.Rect]] = {}
    for edit in spec.get("edits", []):
        target = dict(edit.get("target", {}))
        page_number = int(target.get("page", 0))
        if not 1 <= page_number <= len(doc):
            blocked(f"Invalid PDF page {page_number}.", "Choose an existing one-based page.", "Correct the page? yes/no")
        page = doc[page_number - 1]
        visual_rect = rect_from_target(target, page)
        for existing in by_page.setdefault(page_number, []):
            if existing.intersects(visual_rect):
                blocked(f"PDF visual edits overlap on page {page_number}.", "Merge them into one reviewed region.", "Merge the regions? yes/no")
        by_page[page_number].append(visual_rect)
        raster = image_from_pixmap(page.get_pixmap(matrix=fitz.Matrix(300 / 72, 300 / 72), clip=visual_rect, alpha=False))
        edited = edit_image(raster, edit)
        caption = validate_caption(str(edit.get("caption", "")))
        image_rect = fitz.Rect(visual_rect)
        caption_rect = None
        if caption:
            if edit.get("shrink_inside"):
                caption_height = float(edit.get("caption_height", 42))
                if caption_height <= 16 or caption_height >= visual_rect.height * 0.30:
                    blocked("Caption would shrink the visual by 30% or more.", "Use verified empty space or approve whole-page recomposition.", "Recompose the whole page? yes/no")
                image_rect.y1 -= caption_height
                caption_rect = fitz.Rect(visual_rect.x0, image_rect.y1 + 3, visual_rect.x1, visual_rect.y1)
            else:
                values = edit.get("caption_bbox")
                if not isinstance(values, list) or len(values) != 4:
                    blocked("A PDF caption needs caption_bbox or shrink_inside.", "Provide verified space immediately below the image.", "Shrink the image inside its box? yes/no")
                caption_rect = fitz.Rect(*(float(value) for value in values))
                overlap = max(0, min(visual_rect.x1, caption_rect.x1) - max(visual_rect.x0, caption_rect.x0))
                if caption_rect.y0 < visual_rect.y1 or caption_rect.y0 - visual_rect.y1 > 12 or overlap < visual_rect.width * 0.75:
                    blocked("Caption space is not immediately below the image.", "Choose adjacent same-page space.", "Correct the caption position? yes/no")
                if not page.rect.contains(caption_rect) or not caption_space_blank(page, caption_rect):
                    blocked("Caption space contains source content or leaves the page.", "Shrink the image inside its current box or approve page recomposition.", "Shrink the image inside its box? yes/no")
        stream = io.BytesIO()
        edited.save(stream, format="PNG")
        prepared.append({"page": page_number, "visual_rect": visual_rect, "image_rect": image_rect, "caption_rect": caption_rect, "caption": caption, "image": stream.getvalue()})
    for item in prepared:
        doc[item["page"] - 1].add_redact_annot(item["visual_rect"], fill=(1, 1, 1))
    for page_number in by_page:
        doc[page_number - 1].apply_redactions()
    for item in prepared:
        page = doc[item["page"] - 1]
        page.insert_image(item["image_rect"], stream=item["image"], keep_proportion=True, overlay=True)
        if item["caption_rect"]:
            page.draw_rect(item["caption_rect"], color=None, fill=(1, 1, 1), overlay=True)
            remaining = page.insert_textbox(item["caption_rect"], item["caption"], fontsize=8, fontname="helv", color=(0, 0, 0), align=0, overlay=True)
            if remaining < 0:
                blocked("Caption does not fit its same-page box.", "Shorten the caption or enlarge approved space.", "Shorten the caption? yes/no")
    doc.save(output, garbage=4, deflate=True)
    doc.close()


def insert_paragraph_after(paragraph: Paragraph) -> Paragraph:
    new_element = OxmlElement("w:p")
    paragraph._p.addnext(new_element)
    return Paragraph(new_element, paragraph._parent)


def encode_image(image: Image.Image, member: str) -> bytes:
    stream = io.BytesIO()
    suffix = Path(member).suffix.lower()
    if suffix in (".jpg", ".jpeg"):
        image.convert("RGB").save(stream, format="JPEG", quality=95, subsampling=0)
    else:
        image.save(stream, format="PNG")
    return stream.getvalue()


def apply_docx(source: Path, spec: dict, output: Path) -> None:
    edits = list(spec.get("edits", []))
    replacements: dict[str, bytes] = {}
    with zipfile.ZipFile(source) as archive:
        media_names = set(archive.namelist())
        grouped: dict[str, list[dict]] = {}
        for edit in edits:
            media = str(edit.get("target", {}).get("media", "")).replace("\\", "/")
            if media not in media_names:
                blocked(f"DOCX media target not found: {media}", "Use a media path from the audit report.", "Correct the media target? yes/no")
            grouped.setdefault(media, []).append(edit)
        for media, media_edits in grouped.items():
            image = Image.open(io.BytesIO(archive.read(media))).convert("RGB")
            for edit in media_edits:
                image = edit_image(image, edit)
            replacements[media] = encode_image(image, media)
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as temporary:
            temporary_path = Path(temporary.name)
        try:
            with zipfile.ZipFile(temporary_path, "w", zipfile.ZIP_DEFLATED) as destination:
                for info in archive.infolist():
                    destination.writestr(info, replacements.get(info.filename, archive.read(info.filename)))
            document = Document(temporary_path)
            caption_edits = [edit for edit in edits if validate_caption(str(edit.get("caption", "")))]
            for edit in sorted(caption_edits, key=lambda item: int(item.get("target", {}).get("paragraph", 0)), reverse=True):
                paragraph_number = int(edit.get("target", {}).get("paragraph", 0))
                if not 1 <= paragraph_number <= len(document.paragraphs):
                    blocked(f"DOCX paragraph {paragraph_number} does not exist.", "Use an occurrence from the audit report.", "Correct the paragraph? yes/no")
                image_paragraph = document.paragraphs[paragraph_number - 1]
                if not image_paragraph._p.xpath(".//a:blip"):
                    blocked(f"DOCX paragraph {paragraph_number} contains no inline image.", "Choose the image occurrence paragraph.", "Correct the paragraph? yes/no")
                image_paragraph.paragraph_format.keep_with_next = True
                caption_paragraph = insert_paragraph_after(image_paragraph)
                if "Caption" in [style.name for style in document.styles]:
                    caption_paragraph.style = "Caption"
                caption_paragraph.paragraph_format.keep_together = True
                caption_paragraph.paragraph_format.space_before = Pt(0)
                caption_paragraph.paragraph_format.space_after = Pt(6)
                run = caption_paragraph.add_run(validate_caption(str(edit.get("caption", ""))))
                run.italic = True
                run.font.size = Pt(8)
            document.save(output)
        finally:
            temporary_path.unlink(missing_ok=True)


def wrap_image_pdf(source: Path, edit: dict, output: Path) -> None:
    image = edit_image(Image.open(source).convert("RGB"), edit)
    caption = validate_caption(str(edit.get("caption", "")))
    page_size = landscape(LETTER) if image.width > image.height * 1.25 else LETTER
    width, height = page_size
    margin = 54
    caption_height = 54 if caption else 0
    available_width = width - margin * 2
    available_height = height - margin * 2 - caption_height
    scale = min(available_width / image.width, available_height / image.height)
    draw_width, draw_height = image.width * scale, image.height * scale
    x = (width - draw_width) / 2
    y = margin + caption_height + (available_height - draw_height) / 2
    stream = io.BytesIO()
    image.save(stream, format="PNG")
    c = canvas.Canvas(str(output), pagesize=page_size)
    c.drawImage(ImageReader(io.BytesIO(stream.getvalue())), x, y, draw_width, draw_height, preserveAspectRatio=True, mask="auto")
    if caption:
        c.setFont("Helvetica", 8)
        words = caption.split()
        lines, current = [], ""
        for word in words:
            trial = f"{current} {word}".strip()
            if c.stringWidth(trial, "Helvetica", 8) <= available_width:
                current = trial
            else:
                lines.append(current)
                current = word
        if current:
            lines.append(current)
        caption_y = y - 12
        for line in lines[:4]:
            c.drawString(x, caption_y, line)
            caption_y -= 10
    c.showPage()
    c.save()


def apply(source: Path, spec: dict, output: Path) -> None:
    if output.exists():
        blocked(f"Output already exists: {output}", "Choose a new output path.", "Use a new filename? yes/no")
    output.parent.mkdir(parents=True, exist_ok=True)
    suffix = source.suffix.lower()
    if suffix == ".pdf":
        if output.suffix.lower() != ".pdf":
            blocked("PDF input requires PDF intermediate output.", "Use a .pdf output path.", "Correct the output extension? yes/no")
        apply_pdf(source, spec, output)
    elif suffix == ".docx":
        if output.suffix.lower() != ".docx":
            blocked("DOCX input requires DOCX intermediate output.", "Use a .docx output path.", "Correct the output extension? yes/no")
        apply_docx(source, spec, output)
    elif suffix in (".png", ".jpg", ".jpeg"):
        if output.suffix.lower() != ".pdf":
            blocked("Standalone images require PDF intermediate output.", "Use a .pdf output path.", "Correct the output extension? yes/no")
        edits = list(spec.get("edits", []))
        if len(edits) != 1 or not edits[0].get("target", {}).get("whole_image"):
            blocked("Standalone image spec needs exactly one whole_image edit.", "Set target.whole_image to true.", "Correct the image target? yes/no")
        wrap_image_pdf(source, edits[0], output)
    else:
        blocked(f"Unsupported input format: {suffix or 'none'}.", "Use PDF, DOCX, PNG, or JPG.", "Choose a supported input? yes/no")


def render_pdf(source: Path, output_dir: Path, dpi: int = 144) -> int:
    if output_dir.exists() and any(output_dir.iterdir()):
        blocked(f"Render directory is not empty: {output_dir}", "Choose an empty directory.", "Use a new render directory? yes/no")
    output_dir.mkdir(parents=True, exist_ok=True)
    doc = fitz.open(source)
    matrix = fitz.Matrix(dpi / 72, dpi / 72)
    for index, page in enumerate(doc):
        color = image_from_pixmap(page.get_pixmap(matrix=matrix, alpha=False))
        color.save(output_dir / f"page-{index + 1:04d}-color.png")
        ImageOps.grayscale(color).save(output_dir / f"page-{index + 1:04d}-gray.png")
        laser_preview(color).save(output_dir / f"page-{index + 1:04d}-laser.png")
    count = len(doc)
    doc.close()
    return count


def self_test() -> None:
    with tempfile.TemporaryDirectory(prefix="bw-printable-test-") as temp_dir:
        temp = Path(temp_dir)
        chart = Image.new("RGB", (480, 260), "white")
        draw = ImageDraw.Draw(chart)
        draw.rectangle((20, 20, 225, 220), fill=(200, 0, 0))
        draw.rectangle((255, 20, 460, 220), fill=(0, 100, 0))
        draw.text((50, 230), "A", fill="black", font=load_font(14))
        draw.text((350, 230), "B", fill="black", font=load_font(14))
        report = audit_image(chart)
        assert report["conflicts"], report
        assert any(item["relationship"] == "co-occurring" for item in report["conflicts"]), report
        edited = edit_image(chart, {
            "patterns": [{"color": "#C80000", "tolerance": 8, "style": "diagonal", "spacing": 12}],
            "outlines": [{"color": "#006400", "tolerance": 8, "width": 3}],
            "labels": [{"x": 0.08, "y": 0.08, "text": "A", "size": 14}],
        })
        assert edited.tobytes() != chart.tobytes()
        source = temp / "chart.png"
        chart.save(source)
        output = temp / "chart.pdf"
        apply(source, {"edits": [{
            "target": {"whole_image": True},
            "patterns": [{"color": "#C80000", "tolerance": 8, "style": "diagonal", "spacing": 12}],
            "caption": "The chart compares categories A and B. Category A uses diagonal hatching so it remains distinguishable in grayscale.",
        }]}, output)
        qa = temp / "qa"
        assert render_pdf(output, qa, 96) == 1
        assert (qa / "page-0001-laser.png").exists()
        print("self-test passed")


def parser() -> argparse.ArgumentParser:
    root = argparse.ArgumentParser(description=__doc__)
    sub = root.add_subparsers(dest="command", required=True)
    audit_cmd = sub.add_parser("audit")
    audit_cmd.add_argument("input", type=Path)
    audit_cmd.add_argument("--work-dir", type=Path, required=True)
    audit_cmd.add_argument("--report", type=Path, required=True)
    apply_cmd = sub.add_parser("apply")
    apply_cmd.add_argument("input", type=Path)
    apply_cmd.add_argument("--spec", type=Path, required=True)
    apply_cmd.add_argument("--output", type=Path, required=True)
    render_cmd = sub.add_parser("render")
    render_cmd.add_argument("input", type=Path)
    render_cmd.add_argument("--output-dir", type=Path, required=True)
    render_cmd.add_argument("--dpi", type=int, default=144)
    sub.add_parser("self-test")
    return root


def main() -> int:
    args = parser().parse_args()
    try:
        if args.command == "audit":
            if not args.input.exists():
                blocked(f"Input does not exist: {args.input}", "Choose an existing source.", "Choose another source? yes/no")
            if args.report.exists():
                blocked(f"Report already exists: {args.report}", "Choose a new report path.", "Use a new report path? yes/no")
            report = audit(args.input, args.work_dir)
            args.report.parent.mkdir(parents=True, exist_ok=True)
            args.report.write_text(json.dumps(report, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
            print(args.report.resolve())
        elif args.command == "apply":
            if not args.input.exists() or not args.spec.exists():
                blocked("Input or edit spec does not exist.", "Choose existing files.", "Correct the paths? yes/no")
            spec = json.loads(args.spec.read_text(encoding="utf-8"))
            apply(args.input, spec, args.output)
            print(args.output.resolve())
        elif args.command == "render":
            if not args.input.exists() or args.input.suffix.lower() != ".pdf":
                blocked("Render input must be an existing PDF.", "Choose the intermediate or final PDF.", "Choose another PDF? yes/no")
            print(f"rendered {render_pdf(args.input, args.output_dir, args.dpi)} pages")
        else:
            self_test()
        return 0
    except BWError as exc:
        print(json.dumps({"status": "blocked", "reason": exc.reason, "resolution": exc.resolution, "prompt": exc.prompt}, ensure_ascii=False), file=sys.stderr)
        return 2
    except json.JSONDecodeError as exc:
        print(json.dumps({"status": "blocked", "reason": f"Edit spec JSON is invalid: {exc}", "resolution": "Correct the JSON syntax.", "prompt": "Correct the spec? yes/no"}), file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
